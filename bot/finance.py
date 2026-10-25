"""Касса и отчёты: выручка по каналам, себестоимость, расходы, прибыль, выгрузки и бэкап."""
import io
import json
import zipfile
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

import texts
from db import Base, Expense, Order, PointsTx, Product, StockMovement, User, aware, now
from services import LOCAL_TZ, get_settings

EXPENSE_KINDS = {
    "delivery": "Доставка / курьер",
    "ads": "Реклама",
    "packaging": "Упаковка",
    "rent": "Аренда / связь",
    "salary": "Выплаты",
    "other": "Прочее",
}
PERIODS = {"today": "Сегодня", "week": "7 дней", "month": "30 дней", "all": "Всё время"}


def period_bounds(period: str) -> tuple[datetime | None, datetime]:
    end = now()
    local_midnight = datetime.now(LOCAL_TZ).replace(hour=0, minute=0, second=0, microsecond=0)
    if period == "today":
        return local_midnight, end
    if period == "week":
        return local_midnight - timedelta(days=6), end
    if period == "month":
        return local_midnight - timedelta(days=29), end
    return None, end


def _in(dt, start) -> bool:
    return dt is not None and (start is None or aware(dt) >= start)


def order_cogs(order: Order) -> tuple[int, bool]:
    """Себестоимость товаров в заказе и флаг «у части товаров не указана закупка»."""
    total, unknown = 0, False
    for item in order.items:
        if item.get("cost") is None:
            unknown = True
        else:
            total += item["cost"] * item["qty"]
    return total, unknown


async def report(session: AsyncSession, period: str) -> dict:
    start, _ = period_bounds(period)
    orders = [o for o in (await session.execute(select(Order).where(Order.paid_at.is_not(None)))).scalars()
              if o.status != "cancelled" and _in(o.paid_at, start)]
    receipts = [m for m in (await session.execute(
        select(StockMovement).where(StockMovement.kind == "receipt", StockMovement.cost.is_not(None))
    )).scalars() if _in(m.created_at, start)]
    expenses = [e for e in (await session.execute(select(Expense).order_by(Expense.id.desc()))).scalars()
                if _in(e.created_at, start)]

    channels: dict[str, int] = {}
    revenue = delivery = cogs = 0
    unknown_cost = 0
    for o in orders:
        channels[o.paid_via or "—"] = channels.get(o.paid_via or "—", 0) + o.total
        revenue += o.total
        delivery += o.delivery_price or 0
        c, unknown = order_cogs(o)
        cogs += c
        unknown_cost += unknown
    by_kind: dict[str, int] = {}
    for e in expenses:
        by_kind[e.kind] = by_kind.get(e.kind, 0) + e.amount
    expenses_total = sum(by_kind.values())
    purchases = sum(m.delta * m.cost for m in receipts)

    # Сколько денег лежит в товаре на складе (по средней закупке)
    products = list((await session.execute(select(Product))).scalars())
    all_receipts = list((await session.execute(
        select(StockMovement).where(StockMovement.kind == "receipt", StockMovement.cost.is_not(None))
    )).scalars())
    cost_map: dict[int, list[int]] = {}
    for m in all_receipts:
        acc = cost_map.setdefault(m.product_id, [0, 0])
        acc[0] += m.delta * m.cost
        acc[1] += m.delta
    stock_value = 0
    for p in products:
        acc = cost_map.get(p.id)
        if acc and acc[1]:
            stock_value += sum(s["stock"] for s in p.sizes) * acc[0] // acc[1]

    goods_revenue = revenue - delivery
    return {
        "period": period,
        "period_label": PERIODS.get(period, period),
        "orders": len(orders),
        "revenue": revenue,
        "delivery": delivery,
        "goods_revenue": goods_revenue,
        "cogs": cogs,
        "unknown_cost_orders": unknown_cost,
        "expenses": expenses_total,
        "expenses_by_kind": [{"kind": k, "label": EXPENSE_KINDS.get(k, k), "amount": v} for k, v in by_kind.items()],
        "profit": goods_revenue - cogs - expenses_total,
        "avg_check": revenue // len(orders) if orders else 0,
        "channels": [{"label": k, "amount": v} for k, v in sorted(channels.items(), key=lambda x: -x[1])],
        "purchases": purchases,
        "stock_value": stock_value,
        "expense_list": [expense_json(e) for e in expenses[:50]],
    }


def expense_json(e: Expense) -> dict:
    return {"id": e.id, "kind": e.kind, "label": EXPENSE_KINDS.get(e.kind, e.kind), "amount": e.amount,
            "note": e.note, "admin_name": e.admin_name, "created_at": aware(e.created_at).isoformat()}


async def daily_report_text(session: AsyncSession) -> str:
    settings = await get_settings(session)
    cur = settings["currency"]
    today = await report(session, "today")
    start, _ = period_bounds("today")
    all_orders = list((await session.execute(select(Order))).scalars())
    new_today = [o for o in all_orders if _in(o.created_at, start)]
    count = lambda s: sum(1 for o in all_orders if o.status == s)
    low = []
    for p in (await session.execute(select(Product).where(Product.is_active.is_(True)))).scalars():
        for s in p.sizes:
            if s["stock"] <= settings["low_stock"]:
                low.append(f"{p.title} {s['size']}: {s['stock']}")
    lines = [
        f"📊 <b>Отчёт за {datetime.now(LOCAL_TZ).strftime('%d.%m')}</b>",
        "",
        f"🧾 Новых заказов: <b>{len(new_today)}</b>",
        f"💰 Оплачено: <b>{texts.money(today['revenue'], cur)}</b> ({today['orders']} шт.)",
    ]
    for ch in today["channels"]:
        lines.append(f"   {texts.e(ch['label'])}: {texts.money(ch['amount'], cur)}")
    lines += [
        f"📈 Прибыль за день: <b>{texts.money(today['profit'], cur)}</b>",
        "",
        f"🔎 Проверить оплату: {count('checking')}",
        f"⏳ Ждут оплату: {count('awaiting_payment')}",
        f"🤝 Наличные, ждут выдачи: {count('confirmed')}",
        f"📦 Нужно отправить: {count('paid')}",
        f"🚚 В пути: {count('shipped')}",
    ]
    if low:
        lines += ["", "⚠️ <b>Заканчивается:</b>", *[f"   {texts.e(x)}" for x in low[:15]]]
    return "\n".join(lines)


# ---------- выгрузки в Excel ----------

HEADER_FILL = PatternFill("solid", fgColor="111111")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(wb: Workbook, title: str, headers: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, 1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
    ws.freeze_panes = "A2"


def _local(dt) -> str:
    return aware(dt).astimezone(LOCAL_TZ).strftime("%d.%m.%Y %H:%M") if dt else ""


async def export_xlsx(session: AsyncSession, kind: str) -> tuple[bytes, str]:
    wb = Workbook()
    wb.remove(wb.active)
    if kind in ("orders", "all"):
        orders = list((await session.execute(select(Order).order_by(Order.id.desc()))).scalars())
        _sheet(wb, "Заказы",
               ["№", "Дата", "Статус", "Покупатель", "Телефон", "Telegram", "Товары", "Сумма товаров", "Скидка",
                "Промокод", "Доставка", "Цена доставки", "Итого", "Оплата", "Куда пришли деньги", "Оплачен",
                "Себестоимость", "Прибыль", "Трек", "Ведёт", "Оценка"],
               [[o.id, _local(o.created_at), texts.STATUS_PLAIN.get(o.status, o.status), o.customer_name, o.phone,
                 f"@{o.username}" if o.username else o.user_id,
                 "; ".join(f"{i['title']} {i['size']}×{i['qty']}" for i in o.items),
                 o.subtotal if o.subtotal is not None else o.total, o.discount or 0, o.coupon_code or "",
                 o.delivery, o.delivery_price or 0, o.total,
                 "Наличные" if o.payment_method == "cash" else "Карта", o.paid_via or "", _local(o.paid_at),
                 order_cogs(o)[0], (o.total - (o.delivery_price or 0) - order_cogs(o)[0]) if o.paid_at else "",
                 o.track or "", o.assignee_name or "", o.rating or ""] for o in orders])
    if kind in ("stock", "all"):
        products = list((await session.execute(select(Product).order_by(Product.id))).scalars())
        _sheet(wb, "Остатки", ["Товар", "Категория", "Цена", "Размер", "Остаток"],
               [[p.title, p.category, p.price, s["size"], s["stock"]] for p in products for s in p.sizes])
        moves = list((await session.execute(select(StockMovement).order_by(StockMovement.id.desc()))).scalars())
        _sheet(wb, "Движения склада", ["Дата", "Тип", "Товар", "Размер", "Количество", "Цена закупки", "Комментарий"],
               [[_local(m.created_at), texts.KIND_PLAIN.get(m.kind, m.kind), m.product_title, m.size, m.delta,
                 m.cost or "", m.note] for m in moves])
    if kind in ("finance", "all"):
        expenses = list((await session.execute(select(Expense).order_by(Expense.id.desc()))).scalars())
        _sheet(wb, "Расходы", ["Дата", "Статья", "Сумма", "Комментарий", "Кто внёс"],
               [[_local(e.created_at), EXPENSE_KINDS.get(e.kind, e.kind), e.amount, e.note, e.admin_name or ""]
                for e in expenses])
        for period in ("today", "week", "month", "all"):
            r = await report(session, period)
            ws_rows = [["Заказов оплачено", r["orders"]], ["Выручка", r["revenue"]], ["в т.ч. доставка", r["delivery"]],
                       ["Себестоимость", r["cogs"]], ["Расходы", r["expenses"]], ["Прибыль", r["profit"]],
                       ["Средний чек", r["avg_check"]], ["Закуплено товара", r["purchases"]]]
            ws_rows += [[f"Пришло: {c['label']}", c["amount"]] for c in r["channels"]]
            _sheet(wb, f"Касса · {r['period_label']}", ["Показатель", "Сумма"], ws_rows)
    if kind in ("customers", "all"):
        users = list((await session.execute(select(User).order_by(User.created_at.desc()))).scalars())
        orders = list((await session.execute(select(Order).where(Order.status != "cancelled"))).scalars())
        stats: dict[int, list[int]] = {}
        for o in orders:
            acc = stats.setdefault(o.user_id, [0, 0])
            acc[0] += 1
            acc[1] += o.total if o.paid_at else 0
        _sheet(wb, "Покупатели",
               ["ID", "Telegram", "Имя", "Телефон", "Регистрация", "Баллы", "Заказов", "Потратил", "Не выкупил", "Заметка"],
               [[u.id, f"@{u.username}" if u.username else "", u.full_name or u.first_name or "", u.phone or "",
                 _local(u.registered_at), u.points or 0, stats.get(u.id, [0, 0])[0], stats.get(u.id, [0, 0])[1],
                 "да" if u.no_cash else "", u.note or ""] for u in users])
    if not wb.sheetnames:
        raise ValueError("unknown export")
    buffer = io.BytesIO()
    wb.save(buffer)
    stamp = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
    names = {"orders": "заказы", "stock": "склад", "finance": "касса", "customers": "покупатели", "all": "всё"}
    return buffer.getvalue(), f"dayzdrop-{names.get(kind, kind)}-{stamp}.xlsx"


async def backup_zip(session: AsyncSession) -> tuple[bytes, str]:
    """Полная копия базы в JSON (по таблице на файл), упакованная в zip."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for table in Base.metadata.sorted_tables:
            rows = [dict(r._mapping) for r in (await session.execute(table.select())).all()]
            zf.writestr(f"{table.name}.json", json.dumps(rows, ensure_ascii=False, default=str, indent=1))
    stamp = datetime.now(LOCAL_TZ).strftime("%Y-%m-%d")
    return buffer.getvalue(), f"dayzdrop-backup-{stamp}.zip"


async def points_history(session: AsyncSession, user_id: int, limit: int = 50) -> list[dict]:
    rows = (await session.execute(
        select(PointsTx).where(PointsTx.user_id == user_id).order_by(PointsTx.id.desc()).limit(limit)
    )).scalars()
    return [{"id": t.id, "delta": t.delta, "reason": t.reason, "admin_name": t.admin_name,
             "created_at": aware(t.created_at).isoformat()} for t in rows]
